#!/usr/bin/env python3
"""Build video/slides.pptx — an editable PowerPoint mirror of video/slides.html.

Same nine slides, same headlines, same bullets, same state-machine diagram and
the same testnet hashes. It is not pixel-identical to the HTML deck (PowerPoint
has no CSS), but it carries the same dark theme: near-black background, Hyper
Turquoise accent for headings, amber reserved for warnings, monospace for
anything glanceable (tx hashes, state names, wallet roles).

Deliberately limited to features that survive a round-trip into Canva's PPTX
import: text boxes, autoshapes, pictures, solid/gradient fills, standard fonts.
No SmartArt, no charts, no embedded media, no custom geometry.

Run with the local venv:
    video/.venv-pptx/bin/python video/generate_pptx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
SCREENS = HERE.parent / "submission" / "screens"
OUT = HERE / "slides.pptx"

# --- geometry ---------------------------------------------------------------
# The HTML deck is authored on a 1280x720 stage. A 16:9 PowerPoint slide is
# exactly 12192000 x 6858000 EMU, i.e. exactly 9525 EMU per stage pixel, so the
# HTML coordinates carry over 1:1 and CSS px map to pt at 0.75.
EMU_PER_PX = 9525
STAGE_W, STAGE_H = 1280, 720
PAD_L, PAD_T, PAD_R, PAD_B = 56, 52, 56, 44
CONTENT_W = STAGE_W - PAD_L - PAD_R  # 1168


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


def pt(v: float) -> Pt:
    return Pt(v * 0.75)


# --- palette (mirrors :root in slides.html) ---------------------------------
BG = "0A0C10"
PANEL = "10131A"
FG = "E7ECF2"
MUTED = "94A3B8"
DIM = "64748B"
ACCENT_200 = "A6F7FF"
ACCENT_300 = "66F2FF"
ACCENT_500 = "00F0FF"
AMBER = "FBBF24"
AMBER_200 = "FDE68A"
EMERALD = "34D399"
WHITE_ISH = "F8FAFC"
SLATE_200 = "E2E8F0"
SLATE_300 = "CBD5E1"
CYAN_50 = "CFFAFE"
AMBER_100 = "FEF3C7"
GREEN_100 = "D1FAE5"

SANS = "Arial"
MONO = "Consolas"

# CSS borders/backgrounds are translucent; PowerPoint lines are not. Flatten
# them against the slide background so the deck reads the same.
def blend(hex_fg: str, alpha: float, hex_bg: str = BG) -> str:
    f = [int(hex_fg[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_bg[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(f[i] * alpha + b[i] * (1 - alpha)):02X}" for i in range(3))


LINE = blend("FFFFFF", 0.15)
PANEL_LINE = blend("FFFFFF", 0.15, PANEL)
AMBER_LINE = blend(AMBER, 0.28)
AMBER_FILL = blend(AMBER, 0.08)
ACCENT_LINE = blend("22E7FA", 0.3)
ACCENT_FILL = blend("22E7FA", 0.08)
GREEN_LINE = blend(EMERALD, 0.28)
GREEN_FILL = blend(EMERALD, 0.08)
TAG_FILL = blend("FFFFFF", 0.035)
HOT_LINE = blend(AMBER, 0.45)
HOT_FILL = blend(AMBER, 0.07)
INK = blend("000000", 0.4)

# --- font metrics for the overflow check ------------------------------------
# Liberation Sans is metrically identical to Arial, so sans measurements are
# exact. DejaVu Sans Mono is wider per glyph than Consolas (0.602em vs 0.55em),
# which makes the monospace check conservative — it over-estimates, never under.
FONT_FILES = {
    (SANS, False): "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    (SANS, True): "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    (MONO, False): "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    (MONO, True): "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf",
}
_MEASURE_SCALE = 8
_font_cache: dict = {}


def _font(name: str, bold: bool, size_px: float):
    key = (name, bold, round(size_px, 2))
    if key not in _font_cache:
        path = FONT_FILES[(name, bold)]
        _font_cache[key] = ImageFont.truetype(path, int(round(size_px * _MEASURE_SCALE)))
    return _font_cache[key]


def _text_w(text: str, name: str, bold: bool, size_px: float) -> float:
    return _font(name, bold, size_px).getlength(text) / _MEASURE_SCALE


def wrapped_lines(text: str, name: str, bold: bool, size_px: float, width_px: float) -> int:
    """Number of lines `text` takes when wrapped to width_px (word wrap)."""
    total = 0
    for hard_line in text.split("\n"):
        words = hard_line.split(" ")
        if not words:
            total += 1
            continue
        lines, cur = 1, ""
        for word in words:
            trial = word if not cur else cur + " " + word
            if _text_w(trial, name, bold, size_px) <= width_px or not cur:
                cur = trial
            else:
                lines += 1
                cur = word
        total += lines
    return total


ISSUES: list[str] = []


# --- drawing helpers --------------------------------------------------------
def _solid(fill, hex_color: str):
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hex_color)


def rect(
    slide,
    x, y, w, h,
    fill=None,
    line=None,
    line_w=1.0,
    radius=None,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    sh = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius / min(w, h)
    if fill:
        _solid(sh.fill, fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = pt(line_w)
    else:
        sh.line.fill.background()
    sh.text_frame.text = ""
    _check_bounds(x, y, w, h, "shape")
    return sh


def card(slide, x, y, w, h, fill=PANEL, line=PANEL_LINE, radius=10):
    return rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)


def _check_bounds(x, y, w, h, what):
    if x < 0 or y < 0 or x + w > STAGE_W or y + h > STAGE_H:
        ISSUES.append(f"{what} out of slide bounds: x={x} y={y} w={w} h={h}")


def tbox(
    slide,
    x, y, w, h,
    paras,
    size=15,
    color=MUTED,
    font=SANS,
    bold=False,
    line_spacing=1.45,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    space_after=0,
    hanging=0,
    label="text",
):
    """Add a fixed-size text box.

    `paras` is a string, or a list where each item is a string or a list of
    (text, style-dict) runs. Height usage is measured and checked afterwards.
    """
    if isinstance(paras, str):
        paras = [paras]

    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    used = 0.0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if hanging:
            # Hanging indent so wrapped bullet lines align under the text, not
            # under the bullet. marL/indent are core DrawingML attributes.
            pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
            pPr.set("marL", str(int(px(hanging))))
            pPr.set("indent", str(-int(px(hanging))))
        runs = [(para, {})] if isinstance(para, str) else para

        p_size = size
        p_font = font
        p_any_bold = bold
        plain = ""
        for text, style in runs:
            if "\n" in text:
                raise ValueError(
                    f"tbox[{label}]: newline inside a run is not a PowerPoint line "
                    "break — pass separate paragraphs instead"
                )
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = style.get("font", font)
            f.size = pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.italic = style.get("italic", False)
            f.color.rgb = RGBColor.from_string(style.get("color", color))
            plain += text
            p_size = max(p_size, style.get("size", size))
            if style.get("bold", bold):
                p_any_bold = True
            if style.get("font", font) == MONO:
                p_font = MONO

        if i < len(paras) - 1 and space_after:
            p.space_after = pt(space_after)

        n = wrapped_lines(plain, p_font, p_any_bold, p_size, w - hanging)
        used += n * p_size * line_spacing
        if i < len(paras) - 1:
            used += space_after

    _check_bounds(x, y, w, h, f"textbox[{label}]")
    if used > h + 0.6:
        ISSUES.append(
            f"textbox[{label}] overflows: needs {used:.1f}px in {h}px "
            f"(w={w}, first para: {str(paras[0])[:58]!r})"
        )
    return box


def picture(slide, path: Path, x, y, w, h):
    """Place a screenshot, cropping the bottom like CSS object-fit: cover."""
    pic = slide.shapes.add_picture(str(path), px(x), px(y), width=px(w))
    natural_h = pic.height / EMU_PER_PX
    if h < natural_h:
        pic.crop_bottom = 1 - (h / natural_h)
        pic.height = px(h)
    pic.line.color.rgb = RGBColor.from_string(LINE)
    pic.line.width = pt(1)
    _check_bounds(x, y, w, h, f"picture[{path.name}]")
    return pic


def tag(slide, x, y, text, dot=ACCENT_500, w=None, h=28, size=11):
    """The deck's little mono pill with a glowing status dot."""
    pad = 11
    text_w = _text_w(text, MONO, False, size) * 1.10  # letter-spacing: 0.06em
    if w is None:
        w = pad + 6 + 8 + text_w + pad
    rect(slide, x, y, w, h, fill=TAG_FILL, line=LINE, radius=6)
    rect(slide, x + pad, y + h / 2 - 3, 6, 6, fill=dot, line=None, shape=MSO_SHAPE.OVAL)
    tbox(
        slide, x + pad + 6 + 8, y, w - pad * 2 - 14, h,
        text.upper(), size=size, color=SLATE_300, font=MONO,
        line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE, label="tag",
    )
    return w


# --- deck shell -------------------------------------------------------------
prs = Presentation()
prs.slide_width = px(STAGE_W)
prs.slide_height = px(STAGE_H)
BLANK = prs.slide_layouts[6]
TOTAL = 9


def new_slide(index: int, foot_left: str | None = None, foot_right: str | None = None):
    slide = prs.slides.add_slide(BLANK)

    fill = slide.background.fill
    fill.gradient()
    fill.gradient_angle = 90.0
    stops = fill.gradient_stops
    stops[0].color.rgb = RGBColor.from_string("0C1720")  # cyan-tinted glow, top
    stops[0].position = 0.0
    stops[1].color.rgb = RGBColor.from_string(BG)
    stops[1].position = 0.62

    # progress bar + slide counter, as in the HTML chrome
    rect(slide, 0, 0, STAGE_W * index / TOTAL, 3, fill="22E7FA", line=None,
         shape=MSO_SHAPE.RECTANGLE)
    tbox(
        slide, STAGE_W - PAD_R - 120, STAGE_H - 28, 120, 14,
        f"{index} / {TOTAL}", size=10.5, color=DIM, font=MONO,
        line_spacing=1.0, align=PP_ALIGN.RIGHT, label="counter",
    )
    if foot_left or foot_right:
        y = STAGE_H - PAD_B - 14
        if foot_left:
            tbox(slide, PAD_L, y, CONTENT_W - 300, 14, foot_left.upper(), size=10.5,
                 color=blend("64748B", 0.8), font=MONO, line_spacing=1.0, label="foot-l")
        if foot_right:
            tbox(slide, PAD_L + CONTENT_W - 300, y, 300, 14, foot_right.upper(), size=10.5,
                 color=blend("64748B", 0.8), font=MONO, line_spacing=1.0,
                 align=PP_ALIGN.RIGHT, label="foot-r")
    return slide


def head(slide, eyebrow: str, title: str, kicker: str | None = None) -> float:
    """Eyebrow + h2 (+ kicker). Returns the y where the slide body starts."""
    tbox(slide, PAD_L, PAD_T, CONTENT_W, 16, eyebrow.upper(), size=12, bold=True,
         color=ACCENT_300, font=MONO, line_spacing=1.0, label="eyebrow")
    lines = wrapped_lines(title, SANS, True, 38, CONTENT_W)
    ty = PAD_T + 16 + 10
    th = lines * 38 * 1.1
    tbox(slide, PAD_L, ty, CONTENT_W, th, title, size=38, bold=True, color=WHITE_ISH,
         line_spacing=1.1, label="h2")
    y = ty + th
    if kicker:
        y += 10
        kh = wrapped_lines(kicker, SANS, False, 16, CONTENT_W) * 16 * 1.4
        tbox(slide, PAD_L, y, CONTENT_W, kh, kicker, size=16, color=MUTED,
             line_spacing=1.4, label="kicker")
        y += kh
    return y + 26


def banner(slide, x, y, w, h, runs, kind="amber", size=15):
    line, fill, color = {
        "amber": (AMBER_LINE, AMBER_FILL, AMBER_200),
        "accent": (ACCENT_LINE, ACCENT_FILL, ACCENT_200),
        "green": (GREEN_LINE, GREEN_FILL, "A7F3D0"),
    }[kind]
    rect(slide, x, y, w, h, fill=fill, line=line, radius=10)
    tbox(slide, x + 18, y + 14, w - 36, h - 28, [runs], size=size, color=color,
         line_spacing=1.5, label=f"banner-{kind}")


# ============================================================================
# 1 · TITLE
# ============================================================================
s = new_slide(1)
rect(s, PAD_L, 150, 62, 62, fill=None, line=ACCENT_500, line_w=2.5, shape=MSO_SHAPE.OVAL)
rect(s, PAD_L + 23, 173, 16, 16, fill=ACCENT_500, line=None, shape=MSO_SHAPE.OVAL)
tbox(s, PAD_L, 240, CONTENT_W, 16,
     "Project O · Challenge 1 — Deposit Breakdown & Improvement".upper(),
     size=12, bold=True, color=ACCENT_300, font=MONO, line_spacing=1.0, label="eyebrow")
# Two paragraphs, not one run with "\n": a raw newline inside <a:t> is not a
# line break in PowerPoint.
tbox(s, PAD_L, 266, CONTENT_W, 118, ["Exchange O —", "Deposit Reconciliation Engine"],
     size=54, bold=True, color=WHITE_ISH, line_spacing=1.04, label="h1")
rect(s, PAD_L, 404, 620, 1, fill=blend(ACCENT_500, 0.5), line=None, shape=MSO_SHAPE.RECTANGLE)
tbox(s, PAD_L, 428, 820, 62,
     [[("Closing the gap between ", {}),
       ("“your transaction went through”", {"bold": True, "color": WHITE_ISH}),
       (" and ", {}),
       ("“your collateral is tradable.”", {"bold": True, "color": WHITE_ISH})]],
     size=19, color=MUTED, line_spacing=1.5, label="sub")
tbox(s, PAD_L, 512, 260, 26, "Ryu Matsumoto", size=21, bold=True, color="F1F5F9",
     line_spacing=1.2, label="byline-name")
tbox(s, PAD_L + 216, 517, 420, 18, "Submission · 3-minute walkthrough".upper(),
     size=12, color=DIM, font=MONO, line_spacing=1.0, label="byline-role")
tx = PAD_L
for label, dot in [
    ("Live on Arbitrum Sepolia", EMERALD),
    ("7 real testnet transactions", ACCENT_500),
    ("Working PoC, not mockups", ACCENT_500),
]:
    tx += tag(s, tx, 566, label, dot=dot) + 10

# ============================================================================
# 2 · THE PROBLEM — 8-step journey
# ============================================================================
s = new_slide(
    2,
    "Steps 2 · 3 · 7 can each be a different wallet — nothing labels which is which",
    "submission.md §1",
)
y = head(s, "01 — Breakdown", "The deposit journey today: 8 steps, one blind spot",
         "KOL link → tradable collateral. Every step has an owner and a real completion test.")

STEPS = [
    ("01", "Arrive via KOL link", "attribution can silently drop", "KOL → Project O", False),
    ("02", "Sign in", "method may not match where funds sit", "User + Privy", False),
    ("03", "Identify funding source", "“which wallet has the money?”", "User", False),
    ("04", "Get funds + gas on Arbitrum", "no gas token on the destination chain", "User + bridge / CEX", False),
    ("05", "Approve token spend", "defaults to unlimited allowance", "User + contract", False),
    ("06", "Submit deposit tx → receipt", "receipt ≠ usable funds — the trap", "User + provider", True),
    ("07", "Collateral reflected on Hyperliquid", "real delay, zero visibility today", "External provider", True),
    ("08", "User sees “ready to trade”", "silence reads as failure → re-deposit", "Project O", True),
]
gap, n = 9, 8
cw = (CONTENT_W - gap * (n - 1)) / n
ch = 244
for i, (num, name, fail, who, hot) in enumerate(STEPS):
    x = PAD_L + i * (cw + gap)
    card(s, x, y, cw, ch, fill=HOT_FILL if hot else PANEL,
         line=HOT_LINE if hot else PANEL_LINE, radius=9)
    ix, iw = x + 11, cw - 22
    tbox(s, ix, y + 13, iw, 14, num, size=11, bold=True, color=AMBER if hot else DIM,
         font=MONO, line_spacing=1.0, label=f"step{num}-num")
    tbox(s, ix, y + 35, iw, 54, name, size=13.5, bold=True,
         color=AMBER_100 if hot else SLATE_200, line_spacing=1.25, label=f"step{num}-name")
    tbox(s, ix, y + 95, iw, 92, fail, size=11.5,
         color=blend(AMBER_200, 0.8, HOT_FILL) if hot else blend(MUTED, 0.65),
         line_spacing=1.35, label=f"step{num}-fail")
    tbox(s, ix, y + ch - 25, iw, 14, who.upper(), size=9.5,
         color=blend(AMBER, 0.75, HOT_FILL) if hot else DIM, font=MONO,
         line_spacing=1.0, label=f"step{num}-who")

zy = y + ch + 12
zx = PAD_L + 5 * (cw + gap)
zw = 3 * cw + 2 * gap
rect(s, zx, zy, zw, 32, fill=AMBER_FILL, line=blend(AMBER, 0.4), radius=8)
tbox(s, zx, zy, zw, 32, "The friction zone".upper(), size=11.5, bold=True, color=AMBER_200,
     font=MONO, line_spacing=1.0, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
     label="zone-bar")
banner(
    s, PAD_L, zy + 52, CONTENT_W, 76,
    [("A receipt is ", {}), ("not", {"bold": True, "color": AMBER_100}),
     (" credited collateral. Today steps 6→8 are silent — so “still bridging,” "
      "“no gas,” and “failed” all look identical. Users assume failure and ", {}),
     ("re-deposit", {"bold": True, "color": AMBER_100}), (".", {})],
    kind="amber",
)

# ============================================================================
# 3 · PRIORITISATION — four reasons
# ============================================================================
s = new_slide(3)
y = head(s, "02 — Prioritisation", "Why this is the single biggest friction point",
         "Not one of several equally-good options — four independent reasons.")

REASONS = [
    ("01", "The brief itself names it",
     [("“Complete only when tradable collateral is reflected — ", {}),
      ("not", {"bold": True, "color": WHITE_ISH}),
      (" when a receipt is issued.”", {})]),
    ("02", "Documented in the wild",
     [("Third-party “deposit stuck” guides exist. Their first instruction: ", {}),
      ("“do not send a second deposit.”", {"bold": True, "color": WHITE_ISH})]),
    ("03", "Hits exactly this audience",
     [("Korean CEX withdrawers, KOL-referred, mobile — least equipped to read "
       "ambiguous silence.", {})]),
    ("04", "It causes financial harm",
     [("Duplicate deposits cost real money. Cosmetic friction only costs conversion.", {})]),
]
rw, rh, rgap = (CONTENT_W - 18) / 2, 148, 18
for i, (idx, title, sub) in enumerate(REASONS):
    x = PAD_L + (i % 2) * (rw + rgap)
    ry = y + (i // 2) * (rh + 16)
    card(s, x, ry, rw, rh)
    tbox(s, x + 20, ry + 20, 46, 20, idx, size=15, bold=True, color=ACCENT_300, font=MONO,
         line_spacing=1.0, label=f"reason{idx}-idx")
    tbox(s, x + 66, ry + 18, rw - 86, 24, title, size=17, bold=True, color=SLATE_200,
         line_spacing=1.2, label=f"reason{idx}-title")
    tbox(s, x + 66, ry + 50, rw - 86, rh - 70, [sub], size=14, color=MUTED,
         line_spacing=1.45, label=f"reason{idx}-sub")

banner(
    s, PAD_L, y + 2 * rh + 16 + 18, CONTENT_W, 78,
    [("Considered and rejected as the primary fix: ", {"bold": True, "color": CYAN_50}),
     ("sponsored gas (infra beyond PoC, and the gap remains) · wallet unification "
      "(explicitly out of scope — label, don’t merge) · cutting funnel steps "
      "(no real data — that would be a guess).", {})],
    kind="accent",
)

# ============================================================================
# 4 · THE STATE MACHINE
# ============================================================================
s = new_slide(4)
y = head(s, "03 — The fix", "Deposit Reconciliation Engine",
         "A real backend state machine — persistent records, not a UI illusion.")

PHONE_W4 = 222
LEFT_W = CONTENT_W - 34 - PHONE_W4

NODES = ["SIGNED", "CONFIRMED_ONCHAIN", "BRIDGING", "CREDITED"]
nw, narrow, nh = 200, 30, 44
node_x = []
for i, name in enumerate(NODES):
    x = PAD_L + i * (nw + narrow)
    node_x.append(x)
    done = name == "CREDITED"
    rect(s, x, y, nw, nh, fill=GREEN_FILL if done else PANEL,
         line=GREEN_LINE if done else LINE, radius=8)
    tbox(s, x + 8, y, nw - 16, nh, name, size=13, bold=True,
         color=GREEN_100 if done else SLATE_200, font=MONO, line_spacing=1.0,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label=f"node-{name}")
    if i < len(NODES) - 1:
        tbox(s, x + nw, y, narrow, nh, "→", size=20, color=DIM, line_spacing=1.0,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label="arrow")

EXC = [(0, "STALLED_NO_GAS"), (2, "STALLED_TIMEOUT"), (3, "AMBIGUOUS")]
ey = y + nh + 26
for col, name in EXC:
    x = node_x[col]
    rect(s, x + nw / 2 - 0.5, y + nh, 1, 26, fill=blend(AMBER, 0.35), line=None,
         shape=MSO_SHAPE.RECTANGLE)
    rect(s, x, ey, nw, 36, fill=AMBER_FILL, line=AMBER_LINE, radius=8)
    tbox(s, x + 6, ey, nw - 12, 36, name, size=12, bold=True, color=AMBER_200, font=MONO,
         line_spacing=1.0, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         label=f"exc-{name}")

by = ey + 36 + 26
banner(
    s, PAD_L, by, LEFT_W, 74,
    [("CREDITED needs both sides to agree: ", {"bold": True, "color": GREEN_100}),
     ("the real on-chain receipt read via RPC (never trusting the client) ", {}),
     ("and", {"italic": True}),
     (" the Hyperliquid balance signal.", {})],
    kind="green",
)
banner(
    s, PAD_L, by + 86, LEFT_W, 62,
    [("Disagree past the window → flagged ", {}),
     ("AMBIGUOUS", {"bold": True, "color": AMBER_100}),
     (" for investigation. Never a forever-spinner, never a silent retry.", {})],
    kind="amber",
)
cy = by + 86 + 62 + 12
card(s, PAD_L, cy, LEFT_W, 78)
tbox(
    s, PAD_L + 18, cy + 14, LEFT_W - 36, 50,
    [[("Duplicate-deposit block:", {"bold": True, "color": WHITE_ISH}),
      (" a second deposit for the same wallet + amount while one is in-flight is "
       "refused with a link to the live status. ", {}),
      ("Verified on testnet — HTTP 409.", {"color": ACCENT_300})]],
    size=15, color=MUTED, line_spacing=1.55, label="dup-block",
)

picture(s, SCREENS / "07-status-ambiguous.png", PAD_L + LEFT_W + 34, y, PHONE_W4, 470)
tbox(s, PAD_L + LEFT_W + 34, y + 470 + 9, PHONE_W4, 14,
     "Live status tracker · real app".upper(), size=10.5, color=DIM, font=MONO,
     line_spacing=1.0, align=PP_ALIGN.CENTER, label="shot-cap")

# ============================================================================
# 5 · THE FIVE UX LAYERS
# ============================================================================
s = new_slide(5)
y = head(s, "04 — Five UX layers", "One root cause, five surfaces",
         "Not five unrelated patches — every layer attacks state opacity.")

LAYERS = [
    ("LAYER 01", "Exact-amount approval",
     "Unlimited allowance is opt-in and honestly labelled — never the default."),
    ("LAYER 02", "Plain-language failures",
     "“You don’t have enough ETH to pay the network fee” — zero crypto background assumed."),
    ("LAYER 03", "Full address confirmation",
     "Never truncated, with an explicit lookalike-address warning. Address poisoning is real."),
    ("LAYER 04", "Wallet-role labelling",
     "“Signing in as / Funds coming from / Will be tradable in” — never a bare “your wallet.”"),
    ("LAYER 05", "KOL disclosure",
     "Persistent, compliance-grade: Exchange O is independent — the decision is yours alone."),
]
lgap = 14
lw = (CONTENT_W - lgap * 4) / 5
lh = 228
for i, (num, title, sub) in enumerate(LAYERS):
    x = PAD_L + i * (lw + lgap)
    card(s, x, y, lw, lh)
    rect(s, x + 20, y + 18, 30, 30, fill=ACCENT_FILL, line=ACCENT_LINE, radius=7)
    rect(s, x + 31, y + 29, 8, 8, fill=ACCENT_300, line=None, shape=MSO_SHAPE.OVAL)
    tbox(s, x + 20, y + 60, lw - 40, 14, num, size=10, bold=True, color=ACCENT_300,
         font=MONO, line_spacing=1.0, label=f"{num}-num")
    tbox(s, x + 20, y + 82, lw - 40, 46, title, size=15.5, bold=True, color=SLATE_200,
         line_spacing=1.25, label=f"{num}-title")
    tbox(s, x + 20, y + 132, lw - 40, lh - 150, sub, size=13, color=MUTED,
         line_spacing=1.45, label=f"{num}-sub")

py = y + lh + 24
picture(s, SCREENS / "05-approve.png", PAD_L, py, 172, 206)
picture(s, SCREENS / "04-confirm-address.png", PAD_L + 192, py, 172, 206)
bx = PAD_L + 384
banner(
    s, bx, py + 62, PAD_L + CONTENT_W - bx, 82,
    [("All five ship in the live PoC. The approve screen alone carries three of them — ", {}),
     ("roles, plain-language gas failure, and approval scope", {"bold": True, "color": CYAN_50}),
     (" — on one mobile viewport.", {})],
    kind="accent",
)

# ============================================================================
# 6 · REVIEW-STEP WALKTHROUGH (static screenshots, no live click-through)
# ============================================================================
s = new_slide(6)
y = head(s, "05 — Walkthrough", "The review step, screen by screen",
         "Captured from the deployed PoC on Arbitrum Sepolia — narrated over stills, "
         "not clicked live.")

PHONE_W6, PHONE_H6, PGAP = 216, 460, 22
right_w = PHONE_W6 * 2 + PGAP
left_w6 = CONTENT_W - 34 - right_w

tbox(s, PAD_L, 290, left_w6, 14, "Deployed PoC".upper(), size=11, bold=True, color=DIM,
     line_spacing=1.0, label="deployed-label")
rect(s, PAD_L, 312, left_w6, 54, fill=INK, line=ACCENT_LINE, radius=9)
tbox(s, PAD_L + 18, 312, left_w6 - 36, 54, "https://projecto-blond.vercel.app", size=17,
     color=ACCENT_300, font=MONO, line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE,
     label="url-box")
tbox(
    s, PAD_L, 386, left_w6, 50,
    [[("Real Next.js app, real wallet connection, real Arbitrum Sepolia transactions. ", {}),
      ("No real funds involved.", {"bold": True, "color": WHITE_ISH})]],
    size=15, color=MUTED, line_spacing=1.55, label="demo-note",
)
ty = 448
for label, dot in [
    ("Step 3 — full address, never truncated", ACCENT_500),
    ("Wallet roles: signing in / funding / tradable", ACCENT_500),
    ("Step 4 — approval defaults to this amount only", ACCENT_500),
    ("Plain-language gas failure, with the fix", AMBER),
]:
    tag(s, PAD_L, ty, label, dot=dot)
    ty += 37

rx = PAD_L + left_w6 + 34
for path, cap, offset in [
    ("04-confirm-address.png", "Step 3 · confirm address", 0),
    ("05-approve.png", "Step 4 · approve + deposit", PHONE_W6 + PGAP),
]:
    picture(s, SCREENS / path, rx + offset, y, PHONE_W6, PHONE_H6)
    tbox(s, rx + offset, y + PHONE_H6 + 9, PHONE_W6, 14, cap.upper(), size=10.5, color=DIM,
         font=MONO, line_spacing=1.0, align=PP_ALIGN.CENTER, label="shot-cap")

# ============================================================================
# 7 · REAL TESTNET EVIDENCE
# ============================================================================
s = new_slide(
    7,
    "Mock ERC-20 with identical USDC semantics — Circle’s faucet needs an API key. "
    "Labelled, not hidden.",
    "testnet-evidence.md",
)
y = head(s, "06 — Evidence", "Real transactions. Not simulated.",
         "Every hash below is confirmed on Arbitrum Sepolia and independently verifiable.")

TXS = [
    ("1 — ", "approve(relayer, 25.000000)", " · signed by the user",
     "block 309873172 · success",
     "0xc110d16ae895b7bc9ec8483c6c788a3967f072b16eaaeb0964d46b6f1a3f6022"),
    ("2 — ", "transferFrom(user, relayer, 25.000000)", " · signed by the relayer",
     "block 309873206 · success",
     "0xbaf69d4752b4f1e3a54614e71a1eb25b0c7b553bb829c5a8a5111df1e513e723"),
]
for i, (pre, mono_part, post, meta, hash_hex) in enumerate(TXS):
    cy = y + i * 108
    card(s, PAD_L, cy, CONTENT_W, 96)
    tbox(
        s, PAD_L + 20, cy + 18, CONTENT_W - 40, 20,
        [[(pre, {"bold": True, "color": SLATE_200}),
          (mono_part, {"bold": True, "color": SLATE_200, "font": MONO, "size": 14}),
          (post, {"bold": True, "color": SLATE_200}),
          ("    " + meta, {"color": DIM, "font": MONO, "size": 11})]],
        size=15, line_spacing=1.2, label=f"tx{i}-label",
    )
    rect(s, PAD_L + 20, cy + 46, CONTENT_W - 40, 34, fill=INK, line=PANEL_LINE, radius=8)
    tbox(s, PAD_L + 34, cy + 46, CONTENT_W - 68, 34, hash_hex, size=14, color=SLATE_300,
         font=MONO, line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE, label=f"tx{i}-hash")

vy = y + 216 + 14
tbox(s, PAD_L, vy, CONTENT_W, 18, "verify → https://sepolia.arbiscan.io/tx/<hash>",
     size=13, color=DIM, font=MONO, line_spacing=1.0, label="verify")

STATS = [
    ("7", "real confirmed transactions — none invented"),
    ("0.0", "allowance left after transfer — exact-amount scoping, proven"),
    ("409", "duplicate deposit blocked before any new tx was attempted"),
]
sgap = 14
sw = (CONTENT_W - sgap * 2) / 3
sy = vy + 46
for i, (num, lab) in enumerate(STATS):
    x = PAD_L + i * (sw + sgap)
    card(s, x, sy, sw, 116)
    tbox(s, x + 20, sy + 18, sw - 40, 34, num, size=30, bold=True, color=ACCENT_300,
         font=MONO, line_spacing=1.1, label=f"stat{i}-num")
    tbox(s, x + 20, sy + 60, sw - 40, 44, lab, size=13, color=MUTED, line_spacing=1.4,
         label=f"stat{i}-lab")

# ============================================================================
# 8 · MEASUREMENT
# ============================================================================
s = new_slide(8)
y = head(s, "07 — Measurement", "How we’d know it worked")

card(s, PAD_L, y, CONTENT_W, 158, fill=ACCENT_FILL, line=blend("22E7FA", 0.35))
tbox(s, PAD_L + 26, y + 22, CONTENT_W - 52, 14, "Primary metric".upper(), size=11, bold=True,
     color=blend(ACCENT_300, 0.8, ACCENT_FILL), line_spacing=1.0, label="primary-label")
tbox(
    s, PAD_L + 26, y + 48, CONTENT_W - 52, 76,
    [[("% of deposits reaching ", {}),
      ("CREDITED", {"font": MONO, "size": 24, "color": ACCENT_300}),
      (" without a re-deposit attempt or a support contact", {})]],
    size=27, bold=True, color="F1F5F9", line_spacing=1.3, label="primary-metric",
)
tbox(s, PAD_L, y + 158 + 12, CONTENT_W, 26,
     "Targets the two named failure modes directly — not a generic conversion rate "
     "this fix only influences.",
     size=15, color=MUTED, line_spacing=1.55, label="primary-note")

gy = y + 158 + 12 + 26 + 22
tbox(s, PAD_L, gy, CONTENT_W, 14, "Guardrails — must not regress".upper(), size=11,
     bold=True, color=DIM, line_spacing=1.0, label="guard-label")

GUARDS = [
    ("Time-to-CREDITED p50 / p90",
     "A rising p90 is an infrastructure signal, not a UI one."),
    ("AMBIGUOUS rate",
     "Rising means reconciliation genuinely disagrees more often. Investigate — "
     "don’t extend the timeout."),
    ("Duplicate-block rate",
     "A spike means the status UI isn’t reassuring users fast enough about their "
     "first deposit."),
]
gw = (CONTENT_W - 28) / 3
ghy = gy + 25
for i, (title, sub) in enumerate(GUARDS):
    x = PAD_L + i * (gw + 14)
    card(s, x, ghy, gw, 122)
    tbox(s, x + 20, ghy + 18, gw - 40, 20, title, size=13, bold=True, color=ACCENT_200,
         font=MONO, line_spacing=1.2, label=f"guard{i}-title")
    tbox(s, x + 20, ghy + 46, gw - 40, 62, sub, size=13.5, color=MUTED, line_spacing=1.45,
         label=f"guard{i}-sub")

banner(
    s, PAD_L, ghy + 142, CONTENT_W, 74,
    [("Not available, and not invented: ", {"bold": True, "color": AMBER_100}),
     ("real funnel drop-off, support-ticket volume, and the true production "
      "time-to-credit distribution. The brief withholds them by design.", {})],
    kind="amber",
)

# ============================================================================
# 9 · HONEST CLOSE
# ============================================================================
s = new_slide(9, "Ryu Matsumoto · Project O Challenge 1", "projecto-blond.vercel.app")
y = head(s, "08 — Honest close", "What’s mocked, and what’s next")

PHONE_W9 = 206
left_w9 = CONTENT_W - 34 - PHONE_W9
col_w = (left_w9 - 20) / 2

LIMITS = [
    "No real database — in-memory store; record loss observed live, not theorised",
    "Hyperliquid credit is a timer, not a real balance read",
    "Mock ERC-20, not Circle testnet USDC",
    "Mocked login; no human MetaMask click-through yet",
]
NEXT = [
    "External store — fixes lookups and cross-instance duplicate blocking",
    "Real Hyperliquid balance integration",
    "Re-verify the live deposit rail — the legacy bridge is deprecated",
    "Calibrate AMBIGUOUS thresholds on real bridging latency",
    "Sponsored gas for the no-ETH dead-end",
]
for i, (label, items, col) in enumerate(
    [("Known limitations", LIMITS, AMBER_200), ("Before production", NEXT, SLATE_300)]
):
    x = PAD_L + i * (col_w + 20)
    tbox(s, x, y, col_w, 14, label.upper(), size=11, bold=True,
         color=blend(AMBER, 0.85) if i == 0 else blend("22E7FA", 0.85),
         line_spacing=1.0, label=f"list{i}-label")
    tbox(s, x, y + 27, col_w, 250, ["•  " + t for t in items], size=14.5, color=col,
         line_spacing=1.45, space_after=11, hanging=17, label=f"list{i}")

banner(
    s, PAD_L, y + 296, left_w9, 58,
    [("Every mock is labelled in the UI itself — including the failure screen when "
      "the PoC loses its own record.", {})],
    kind="accent",
)

picture(s, SCREENS / "06-status-credited.png", PAD_L + left_w9 + 34, y, PHONE_W9, 424)
tbox(s, PAD_L + left_w9 + 34, y + 424 + 9, PHONE_W9, 14, "Limitation, surfaced honestly".upper(),
     size=10.5, color=DIM, font=MONO, line_spacing=1.0, align=PP_ALIGN.CENTER,
     label="shot-cap")

# ============================================================================
prs.save(OUT)

if ISSUES:
    print(f"LAYOUT PROBLEMS ({len(ISSUES)}):", file=sys.stderr)
    for msg in ISSUES:
        print("  -", msg, file=sys.stderr)
    sys.exit(1)

print(f"wrote {OUT} — {len(prs.slides._sldIdLst)} slides, layout clean")
